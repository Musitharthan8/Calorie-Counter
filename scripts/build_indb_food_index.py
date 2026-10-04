#!/usr/bin/env python3
"""Build a canonical SQLite nutrition index from Anuvaad INDB.

Input:
  Anuvaad_INDB_2024.11.xlsx from the Anuvaad Indian Nutrient Databank portal.

This importer is intentionally separate from redistribution of the generated database.
The Anuvaad portal describes INDB as open-access and the associated 2024 article is CC BY 4.0,
but the download page itself does not currently display an explicit dataset licence notice.
Therefore --license is REQUIRED: whoever produces a redistributable asset must supply the exact
licence they have independently verified.

Requires:
  pip install openpyxl

Example after licence verification:
  python scripts/build_indb_food_index.py \
      --input Anuvaad_INDB_2024.11.xlsx \
      --dataset-version 2024.11 \
      --license "CC BY 4.0"
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import statistics
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("error: openpyxl missing - run: pip install openpyxl==3.1.5")

SOURCE_URL = "https://www.anuvaad.org.in/indian-nutrient-databank/"
SOURCE_NAME = "Anuvaad Indian Nutrient Databank (INDB)"
SOURCE_KIND = "INDB"
DEFAULT_ATTRIBUTION = "Anuvaad Solutions LLP, Indian Nutrient Databank (INDB)"

# App nutrient key -> (workbook column, unit factor, canonical unit)
# Free sugar is deliberately NOT mapped to total/added sugar; it remains its own concept.
NUTRIENTS = {
    "calories": ("energy_kcal", 1.0, "kcal"),
    "protein": ("protein_g", 1.0, "g"),
    "carbs": ("carb_g", 1.0, "g"),
    "fat": ("fat_g", 1.0, "g"),
    "freeSugar": ("freesugar_g", 1.0, "g"),
    "fiber": ("fibre_g", 1.0, "g"),
    "saturatedFat": ("sfa_mg", 0.001, "g"),
    "monounsaturatedFat": ("mufa_mg", 0.001, "g"),
    "polyunsaturatedFat": ("pufa_mg", 0.001, "g"),
    "cholesterol": ("cholesterol_mg", 1.0, "mg"),
    "sodium": ("sodium_mg", 1.0, "mg"),
    "potassium": ("potassium_mg", 1.0, "mg"),
    "magnesium": ("magnesium_mg", 1.0, "mg"),
    "calcium": ("calcium_mg", 1.0, "mg"),
    "iron": ("iron_mg", 1.0, "mg"),
    "zinc": ("zinc_mg", 1.0, "mg"),
    "phosphorus": ("phosphorus_mg", 1.0, "mg"),
    "vitaminA": ("vita_ug", 1.0, "ug"),
    "vitaminC": ("vitc_mg", 1.0, "mg"),
    "vitaminB6": ("vitb6_mg", 1.0, "mg"),
    "niacin": ("vitb3_mg", 1.0, "mg"),
}

MACRO_KEYS = ("calories", "protein", "carbs", "fat")
ALIAS_RE = re.compile(r"\(([^()]{2,80})\)\s*$")
SPACE_RE = re.compile(r"\s+")


def normalize_name(value: str) -> str:
    value = value.replace("’", "'").strip().lower()
    value = re.sub(r"[^\w']+", " ", value, flags=re.UNICODE)
    return SPACE_RE.sub(" ", value).strip()


def value(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not (result == result) or result in (float("inf"), float("-inf")):
        return None
    return result


def canonical_name_and_aliases(food_name: str) -> tuple[str, list[str]]:
    raw = SPACE_RE.sub(" ", food_name.strip())
    local_match = ALIAS_RE.search(raw)
    without_local = ALIAS_RE.sub("", raw).strip()
    canonical = without_local.split(",", 1)[0].strip() or without_local or raw
    canonical = canonical[:120]

    candidates = [raw, without_local]
    if local_match:
        candidates.append(local_match.group(1).strip())
    first_clause = raw.split(",", 1)[0].strip()
    if first_clause:
        candidates.append(first_clause)

    aliases: list[str] = []
    seen = {normalize_name(canonical)}
    for candidate in candidates:
        candidate = SPACE_RE.sub(" ", candidate.strip())
        normalized = normalize_name(candidate)
        if candidate and normalized and normalized not in seen:
            aliases.append(candidate[:160])
            seen.add(normalized)
    return canonical, aliases


def normalize_serving_unit(raw: str) -> str:
    text = SPACE_RE.sub(" ", raw.strip().lower())
    replacements = {
        "tablespoon": "tbsp",
        "tablespoons": "tbsp",
        "teaspoon": "tsp",
        "teaspoons": "tsp",
        "cups": "cup",
        "pieces": "piece",
        "slices": "slice",
        "bowls": "bowl",
        "plates": "plate",
        "servings": "serving",
        "glasses": "glass",
        "katoris": "katori",
    }
    return replacements.get(text, text[:48] or "serving")


def derived_serving_grams(row: tuple, columns: dict[str, int]) -> float | None:
    """Infer serving grams only when at least two independent nutrient ratios agree.

    INDB includes per-100g and per-serving values but no explicit gram-weight field in the public
    converter contract. Ratios are vulnerable to rounding, so a single ratio is not accepted.
    """
    pairs = (
        ("energy_kcal", "unit_serving_energy_kcal"),
        ("protein_g", "unit_serving_protein_g"),
        ("carb_g", "unit_serving_carb_g"),
        ("fat_g", "unit_serving_fat_g"),
    )
    estimates: list[float] = []
    for base_column, serving_column in pairs:
        if base_column not in columns or serving_column not in columns:
            continue
        base = value(row[columns[base_column]])
        serving = value(row[columns[serving_column]])
        if base is None or serving is None or base <= 0 or serving <= 0:
            continue
        grams = serving / base * 100.0
        if not 5.0 <= grams <= 2500.0:
            return None  # A contradictory ratio must not disappear from the agreement check.
        estimates.append(grams)

    if len(estimates) < 2:
        return None
    median = statistics.median(estimates)
    if median <= 0:
        return None
    max_relative_error = max(abs(candidate - median) / median for candidate in estimates)
    if max_relative_error > 0.10:
        return None
    return round(median, 2)


def create_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA foreign_keys = ON;

        CREATE TABLE foods (
            source_food_id TEXT PRIMARY KEY,
            canonical_name TEXT NOT NULL,
            normalized_name TEXT NOT NULL,
            region TEXT,
            brand TEXT,
            preparation TEXT,
            calories REAL NOT NULL,
            protein REAL NOT NULL,
            carbs REAL NOT NULL,
            fat REAL NOT NULL,
            micronutrients_json TEXT NOT NULL,
            source_url TEXT,
            search_text TEXT NOT NULL
        );

        CREATE TABLE aliases (
            source_food_id TEXT NOT NULL,
            name TEXT NOT NULL,
            normalized_name TEXT NOT NULL,
            language_tag TEXT,
            PRIMARY KEY (source_food_id, normalized_name),
            FOREIGN KEY (source_food_id) REFERENCES foods(source_food_id) ON DELETE CASCADE
        );

        CREATE TABLE portions (
            source_food_id TEXT NOT NULL,
            amount REAL NOT NULL,
            unit TEXT NOT NULL,
            grams REAL NOT NULL,
            description TEXT,
            is_default INTEGER NOT NULL DEFAULT 0,
            is_derived INTEGER NOT NULL DEFAULT 0,
            PRIMARY KEY (source_food_id, unit, amount, grams),
            FOREIGN KEY (source_food_id) REFERENCES foods(source_food_id) ON DELETE CASCADE
        );

        CREATE INDEX foods_normalized_name_idx ON foods(normalized_name);
        CREATE INDEX aliases_normalized_name_idx ON aliases(normalized_name);
        CREATE INDEX portions_food_idx ON portions(source_food_id);
        """
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_manifest(
    db_path: Path,
    *,
    dataset_version: str,
    license_name: str,
    attribution: str,
    record_count: int,
) -> Path:
    manifest = {
        "source": SOURCE_KIND,
        "datasetName": SOURCE_NAME,
        "datasetVersion": dataset_version,
        "sourceUrl": SOURCE_URL,
        "license": license_name,
        "attribution": attribution,
        "generatedAtUtc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "recordCount": record_count,
        "sha256": sha256_file(db_path),
    }
    path = db_path.with_suffix(".manifest.json")
    path.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return path


def build(
    workbook_path: Path,
    output: Path,
    dataset_version: str,
    license_name: str,
    attribution: str,
) -> int:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        workbook.close()
        raise SystemExit("INDB workbook has no rows")

    header = [str(cell).strip() if cell is not None else "" for cell in header_row]
    columns = {name: index for index, name in enumerate(header)}
    required = {
        "food_code",
        "food_name",
        "energy_kcal",
        "protein_g",
        "carb_g",
        "fat_g",
        "servings_unit",
        "unit_serving_energy_kcal",
    }
    missing = sorted(required - columns.keys())
    if missing:
        workbook.close()
        raise SystemExit("INDB workbook is missing columns: " + ", ".join(missing))

    data_rows = list(rows)
    codes = [str(row[columns["food_code"]] or "").strip() for row in data_rows]
    seen = set()
    for code in codes:
        if code and code in seen:
            workbook.close()
            raise ValueError(f"Duplicate INDB food_code: {code}")
        seen.add(code)

    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    conn = sqlite3.connect(output)
    create_schema(conn)

    seen_codes: set[str] = set()
    record_count = 0
    try:
        for row in data_rows:
            code = str(row[columns["food_code"]] or "").strip()
            raw_name = str(row[columns["food_name"]] or "").strip()
            if not code or not raw_name or code in seen_codes:
                continue
            seen_codes.add(code)

            parsed: dict[str, tuple[float, str]] = {}
            for key, (column_name, factor, unit) in NUTRIENTS.items():
                if column_name not in columns:
                    continue
                amount = value(row[columns[column_name]])
                if amount is None or amount < 0:
                    continue
                parsed[key] = (amount * factor, unit)

            d2 = value(row[columns["vitd2_ug"]]) if "vitd2_ug" in columns else None
            d3 = value(row[columns["vitd3_ug"]]) if "vitd3_ug" in columns else None
            if d2 is not None and d3 is not None and d2 >= 0 and d3 >= 0:
                parsed["vitaminD"] = (d2 + d3, "ug")

            if any(key not in parsed for key in MACRO_KEYS):
                continue

            canonical_name, aliases = canonical_name_and_aliases(raw_name)
            macros = {key: parsed[key][0] for key in MACRO_KEYS}
            micronutrients = {
                key: {"amount": amount, "unit": unit}
                for key, (amount, unit) in parsed.items()
                if key not in MACRO_KEYS
            }
            search_text = normalize_name(" ".join([canonical_name, *aliases]))
            source_url = SOURCE_URL
            conn.execute(
                """
                INSERT INTO foods(
                    source_food_id, canonical_name, normalized_name, region, brand, preparation,
                    calories, protein, carbs, fat, micronutrients_json, source_url, search_text
                ) VALUES (?, ?, ?, NULL, NULL, NULL, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    code,
                    canonical_name,
                    normalize_name(canonical_name),
                    macros["calories"],
                    macros["protein"],
                    macros["carbs"],
                    macros["fat"],
                    json.dumps(micronutrients, sort_keys=True, separators=(",", ":")),
                    source_url,
                    search_text,
                ),
            )

            # Preserve full/local/common names as aliases. Local names are not assigned a language
            # tag because the workbook supplies romanized names, not a reliable BCP-47 language.
            for alias in aliases:
                conn.execute(
                    """
                    INSERT OR IGNORE INTO aliases(
                        source_food_id, name, normalized_name, language_tag
                    ) VALUES (?, ?, ?, NULL)
                    """,
                    (code, alias, normalize_name(alias)),
                )

            serving_unit_raw = str(row[columns["servings_unit"]] or "").strip()
            serving_grams = derived_serving_grams(row, columns)
            if serving_unit_raw and serving_grams is not None:
                conn.execute(
                    """
                    INSERT OR IGNORE INTO portions(
                        source_food_id, amount, unit, grams, description, is_default, is_derived
                    ) VALUES (?, 1.0, ?, ?, ?, 1, 1)
                    """,
                    (
                        code,
                        normalize_serving_unit(serving_unit_raw),
                        serving_grams,
                        f"1 {serving_unit_raw}",
                    ),
                )

            record_count += 1

        conn.commit()
        conn.execute("PRAGMA optimize")
    finally:
        conn.close()
        workbook.close()

    write_manifest(
        output,
        dataset_version=dataset_version,
        license_name=license_name,
        attribution=attribution,
        record_count=record_count,
    )
    return record_count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build canonical Anuvaad INDB nutrition index")
    parser.add_argument("--input", required=True, type=Path, help="Anuvaad_INDB_*.xlsx")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(
            "android/app/src/main/assets/nutrition/indb/indb_foods.sqlite"
        ),
    )
    parser.add_argument("--dataset-version", required=True)
    parser.add_argument(
        "--license",
        required=True,
        dest="license_name",
        help="Exact dataset licence independently verified for redistribution",
    )
    parser.add_argument("--attribution", default=DEFAULT_ATTRIBUTION)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    if not input_path.is_file():
        raise SystemExit(f"INDB workbook not found: {input_path}")
    license_name = args.license_name.strip()
    if not license_name:
        raise SystemExit("--license must not be blank")
    dataset_version = args.dataset_version.strip()
    if not dataset_version:
        raise SystemExit("--dataset-version must not be blank")

    output = args.output.resolve()
    count = build(
        workbook_path=input_path,
        output=output,
        dataset_version=dataset_version,
        license_name=license_name,
        attribution=args.attribution.strip() or DEFAULT_ATTRIBUTION,
    )
    print(f"Wrote INDB index: {output} ({count} complete-macro foods)")
    print(f"Wrote manifest: {output.with_suffix('.manifest.json')}")
    print("Reminder: committing the generated asset is a separate licensing decision.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
